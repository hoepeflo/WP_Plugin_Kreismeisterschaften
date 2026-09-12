#!/usr/bin/env python3
"""
Konvertiert den NSSV-Disziplinplan (xlsx) in das Regeltabellen-JSON des KM-Meldeportals
(Format "kmm-regeltabelle", Version 1) und erzeugt einen Prüfbericht.

Aufruf:
    python3 tools/konvertiere-disziplinplan.py \
        docs/quellen/01A1_Disziplinenplan_2026-Aktuell_17.02.xlsx \
        --sportjahr 2026 \
        --json docs/regeltabelle-2026.json \
        --bericht docs/pruefbericht-2026.md \
        --uebersicht docs/regeltabelle-2026-uebersicht.md \
        [--bogen-csv docs/quellen/bogen-2026.csv]

Voraussetzung: Python 3.10+, `pip install openpyxl`.

Das Skript gehört nicht ins ausgelieferte Plugin. Es kennt die Eigenheiten der Datei 2026
(mehrere Tabellen auf einem Blatt, Rotfüllung = kein Startrecht, uneinheitliche Verweise,
Blasrohr unter 1.12 statt 12.10, getauschte Alter/Klasse-Zeilen, geteilte Klassennummern
beim Lichtschießen). Jede Zelle, die nicht eindeutig zugeordnet werden kann, landet im
Prüfbericht mit Koordinate, Rohwert, Füllfarbe und der getroffenen Annahme.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import date
from typing import Any

try:
    import openpyxl
    from openpyxl.utils import column_index_from_string, get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl fehlt: pip install openpyxl")

# ---------------------------------------------------------------------------
# Klassensatz (muss zum Seed in src/Domain/Seed/Klassensatz.php passen)
# ---------------------------------------------------------------------------

KLASSEN: dict[str, dict[str, str]] = {
    # gruppe -> {"<nummer><geschlecht>": bezeichnung}
    "freihand": {
        "20m": "Schüler m", "21w": "Schüler w", "30m": "Jugend m", "31w": "Jugend w",
        "42m": "Junioren II m", "43w": "Junioren II w", "40m": "Junioren I m", "41w": "Junioren I w",
        "10m": "Herren I", "11w": "Damen I", "12m": "Herren II", "13w": "Damen II",
        "14m": "Herren III", "15w": "Damen III", "16m": "Herren IV", "17w": "Damen IV",
        "18m": "Herren V", "19w": "Damen V",
        "40x": "Team Junioren", "10x": "Team Damen/Herren",
    },
    "auflage": {
        "50m": "Senioren 0 m", "51w": "Senioren 0 w", "70m": "Senioren I m", "71w": "Senioren I w",
        "72m": "Senioren II m", "73w": "Senioren II w", "74m": "Senioren III m", "75w": "Senioren III w",
        "76m": "Senioren IV m", "77w": "Senioren IV w", "78m": "Senioren V m", "79w": "Senioren V w",
        "80m": "Senioren VI m", "81w": "Senioren VI w",
    },
    "fitasc": {"68x": "Junioren", "60m": "Herren", "61w": "Damen", "62x": "Senioren", "64x": "Veteranen", "66x": "Master"},
    "lichtschiessen": {
        "26m": "Schüler IV m", "27w": "Schüler IV w", "24m": "Schüler III m", "25w": "Schüler III w",
        "22m": "Schüler II m", "23w": "Schüler II w", "20m": "Schüler I m", "21w": "Schüler I w",
    },
    "blasrohr": {
        "24m": "Schüler III m", "25w": "Schüler III w", "22m": "Schüler II m", "23w": "Schüler II w",
        "20m": "Schüler I m", "21w": "Schüler I w", "30m": "Jugend m", "31w": "Jugend w",
        "42m": "Junioren II m", "43w": "Junioren II w", "40m": "Junioren I m", "41w": "Junioren I w",
        "10m": "Herren I", "11w": "Damen I", "12m": "Herren II", "13w": "Damen II",
        "14m": "Herren III", "15w": "Damen III", "16m": "Herren IV", "17w": "Damen IV",
        "18m": "Herren V", "19w": "Damen V",
    },
    "bogen": {
        "24m": "Schüler C m", "25w": "Schüler C w", "22m": "Schüler B m", "23w": "Schüler B w",
        "20m": "Schüler A m", "21w": "Schüler A w", "30m": "Jugend m", "31w": "Jugend w",
        "40m": "Junioren m", "41w": "Junioren w", "10m": "Herren", "11w": "Damen",
        "12m": "Master m", "13w": "Master w", "14m": "Senioren m", "15w": "Senioren w",
    },
    "para": {"90x": "SH2/AB2 m/w mit HM", "92m": "SH1/AB1 m ohne HM", "93w": "SH1/AB1 w ohne HM", "94x": "AB3 m/w mit HM", "96x": "SH3 m/w ohne HM"},
}

# Nicht im Umfang (Konzept Abschnitt 3): Target Sprint, Sommerbiathlon, Laufende Scheibe, Armbrust.
UEBERSPRINGEN_PREFIX = ("4.", "5.", "8.")

# Ergebnisformat: Zehntelringe nur für Luftgewehr 1.10 (Entscheidung Florian, änderbar im Backend).
ZEHNTEL = {"1.10"}

# Tarif-Überschreibungen je Disziplin (Konzept 11.3).
TARIF_OVERRIDE = {"11.": 2.50}

# Füllfarben.
ROT = {"FFFF0000", "FFD20000"}
GRUEN = {"FF92D050", "FF89C34B"}
GELB = {"FFFFFF00"}
BLAU = {"FF00B0F0", "FF00CCFF"}
ORANGE = {"FFFFC000"}


# ---------------------------------------------------------------------------
# Hilfsklassen
# ---------------------------------------------------------------------------

@dataclass
class Zelle:
    koord: str
    wert: str
    fuellung: str  # rot | gruen | gelb | blau | orange | keine | sonstige:<rgb>


@dataclass
class Befund:
    art: str  # unklar | hinweis | uebersprungen | info
    koord: str
    wert: str
    fuellung: str
    disziplin: str
    klasse: str
    annahme: str


@dataclass
class Regel:
    klasse: str  # "10m" oder "para:92m"
    einzel: str = "keine"
    einzel_ziel: str | None = None
    mannschaft: str = "keine"
    mannschaft_ziel: str | None = None
    mindestalter: int | None = None
    hinweis: str = ""

    def leer(self) -> bool:
        return self.einzel == "keine" and self.mannschaft == "keine"

    def to_dict(self) -> dict[str, Any]:
        d: dict[str, Any] = {"klasse": self.klasse, "einzel": self.einzel, "mannschaft": self.mannschaft}
        if self.einzel_ziel:
            d["einzel_ziel"] = self.einzel_ziel
        if self.mannschaft_ziel:
            d["mannschaft_ziel"] = self.mannschaft_ziel
        if self.mindestalter is not None:
            d["mindestalter"] = self.mindestalter
        if self.hinweis:
            d["hinweis"] = self.hinweis
        return d


@dataclass
class Disziplin:
    kennzahl: str
    bezeichnung: str
    gruppe: str
    typ: str = "normal"
    hinweis: str = ""
    sortierung: int = 0
    quelle: str = ""
    regeln: dict[str, Regel] = field(default_factory=dict)
    hat_mannschaft: bool = False

    def regel(self, klasse: str) -> Regel:
        if klasse not in self.regeln:
            self.regeln[klasse] = Regel(klasse=klasse)
        return self.regeln[klasse]

    def to_dict(self) -> dict[str, Any]:
        groesse = 2 if self.typ == "mixteam" else (0 if self.typ == "bogen" or not self.hat_mannschaft else 3)
        tarif = None
        for prefix, betrag in TARIF_OVERRIDE.items():
            if self.kennzahl.startswith(prefix):
                tarif = betrag
        d = {
            "kennzahl": self.kennzahl,
            "bezeichnung": self.bezeichnung,
            "gruppe": self.gruppe,
            "typ": self.typ,
            "angeboten": True,
            "mannschaft_groesse": groesse,
            "ergebnis_format": "zehntel" if self.kennzahl in ZEHNTEL else "ganz",
            "tarif_override": tarif,
            "mannschaft_startgeld": 0.0,
            "mixteam_kennzahl_modus": "team" if self.typ == "mixteam" else None,
            "hinweis": self.hinweis,
            "sortierung": self.sortierung,
            "regeln": [r.to_dict() for r in self.regeln.values() if not r.leer()],
        }
        return d


class Blatt:
    """Zugriff auf Zellen mit Auflösung verbundener Bereiche und Füllfarben."""

    def __init__(self, ws) -> None:  # noqa: ANN001
        self.ws = ws
        self.merged_top: dict[tuple[int, int], tuple[int, int]] = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self.merged_top[(r, c)] = (rng.min_row, rng.min_col)

    def zelle(self, row: int, col: int | str) -> Zelle:
        if isinstance(col, str):
            col = column_index_from_string(col)
        top = self.merged_top.get((row, col), (row, col))
        c = self.ws.cell(row=top[0], column=top[1])
        wert = c.value
        wert = "" if wert is None else str(wert).strip()
        return Zelle(koord=f"{get_column_letter(col)}{row}", wert=wert, fuellung=self._fuellung(c))

    def text(self, row: int, col: int | str) -> str:
        return self.zelle(row, col).wert

    @staticmethod
    def _fuellung(c) -> str:  # noqa: ANN001
        f = c.fill
        if not f or f.fill_type != "solid":
            return "keine"
        col = f.fgColor
        if col.type == "rgb" and isinstance(col.rgb, str):
            rgb = col.rgb.upper()
            if rgb in ROT:
                return "rot"
            if rgb in GRUEN:
                return "gruen"
            if rgb in GELB:
                return "gelb"
            if rgb in BLAU:
                return "blau"
            if rgb in ORANGE:
                return "orange"
            if rgb in ("00000000", "FFFFFFFF"):
                return "keine"
            return f"sonstige:{rgb}"
        if col.type == "theme":
            return "keine" if col.theme == 0 else f"theme:{col.theme}"
        if col.type == "indexed":
            return "keine" if col.indexed in (64, 65) else f"indexed:{col.indexed}"
        return "keine"


# ---------------------------------------------------------------------------
# Zellinterpretation
# ---------------------------------------------------------------------------

RE_VERWEIS = re.compile(r"^b\s*\.?\s*(\d{2})\s*\*{0,2}\s*$", re.I)
RE_GESCHLECHT_VERWEIS = re.compile(r"^m\s*b?\s*\.?\s*(\d{2})\s*/\s*w\s*b?\s*\.?\s*(\d{2})\s*$", re.I)
RE_EIGEN = re.compile(r"^(E|M)(\s*\*+)?$")


def interpretiere(z: Zelle) -> tuple[str, Any, str]:
    """Liefert (art, wert, hinweis). art: keine | eigen | verweis | geschlecht_verweis | ab18 |
    team | bei_team | unklar."""
    w = z.wert.replace("\n", " ").strip()
    if w == "":
        if z.fuellung in ("rot", "keine"):
            return "keine", None, ""
        return "unklar", None, ""
    if w.lower() == "x":
        return "keine", None, ""
    m = RE_EIGEN.match(w)
    if m:
        return "eigen", m.group(1), ("Meldung zum DSB anders" if m.group(2) else "")
    m = RE_VERWEIS.match(w)
    if m:
        return "verweis", int(m.group(1)), ""
    m = RE_GESCHLECHT_VERWEIS.match(w)
    if m:
        return "geschlecht_verweis", {"m": int(m.group(1)), "w": int(m.group(2))}, ""
    lw = w.lower().replace(" ", "")
    if lw.startswith("ab18"):
        return "ab18", 18, w
    if lw in ("teamjunioren",):
        return "team", "40x", ""
    if lw in ("teamdamen/herren", "teamherren/damen"):
        return "team", "10x", ""
    if lw in ("beijunioren", "beiteamjunioren"):
        return "bei_team", "40x", ""
    if lw in ("beiteamdamen/herren", "beiteamherren/damen"):
        return "bei_team", "10x", ""
    return "unklar", w, ""


# ---------------------------------------------------------------------------
# Blockdefinitionen (Spalten je Klasse: Einzel-Spalte, Mannschafts-Spalte)
# ---------------------------------------------------------------------------

@dataclass
class Spalte:
    klasse: str          # "10m"
    einzel: str          # Spaltenbuchstabe
    mannschaft: str | None  # Spaltenbuchstabe (None = keine Mannschaftsspalte)
    geteilt: bool = False   # Mannschaftsspalte mit der Partnerklasse geteilt (Schüler/Jugend)


def spalten_hauptblock() -> list[Spalte]:
    return [
        Spalte("20m", "D", "E", True), Spalte("21w", "F", "E", True),
        Spalte("30m", "G", "H", True), Spalte("31w", "I", "H", True),
        Spalte("42m", "J", "K"), Spalte("43w", "L", "M"),
        Spalte("40m", "N", "O"), Spalte("41w", "P", "Q"),
        Spalte("10m", "R", "S"), Spalte("11w", "T", "U"),
        Spalte("12m", "V", "W"), Spalte("13w", "X", "Y"),
        Spalte("14m", "Z", "AA"), Spalte("15w", "AB", "AC"),
        Spalte("16m", "AD", "AE"), Spalte("17w", "AF", "AG"),
        Spalte("18m", "AH", "AI"), Spalte("19w", "AJ", "AK"),
    ]


def spalten_para() -> list[Spalte]:
    return [Spalte("90x", "Z", "AA"), Spalte("92m", "AB", "AC"), Spalte("93w", "AD", "AE"), Spalte("94x", "AF", "AG"), Spalte("96x", "AH", "AI")]


def spalten_auflage() -> list[Spalte]:
    cols = ["D", "F", "H", "J", "L", "N", "P", "R", "T", "V", "X", "Z", "AB", "AD"]
    klassen = ["50m", "51w", "70m", "71w", "72m", "73w", "74m", "75w", "76m", "77w", "78m", "79w", "80m", "81w"]
    out = []
    for k, c in zip(klassen, cols):
        ci = column_index_from_string(c)
        out.append(Spalte(k, c, get_column_letter(ci + 1)))
    return out


def spalten_fitasc() -> list[Spalte]:
    return [Spalte("68x", "D", "F"), Spalte("60m", "H", "I"), Spalte("61w", "J", "K"), Spalte("62x", "L", "N"), Spalte("64x", "P", "R"), Spalte("66x", "T", "V")]


def spalten_blasrohr_erwachsene() -> list[Spalte]:
    klassen = ["10m", "11w", "12m", "13w", "14m", "15w", "16m", "17w", "18m", "19w"]
    out = []
    for i, k in enumerate(klassen):
        ci = 3 + i * 2  # C, E, G, ...
        out.append(Spalte(k, get_column_letter(ci), get_column_letter(ci + 1)))
    return out


def spalten_blasrohr_jugend() -> list[Spalte]:
    return [
        Spalte("24m", "C", "D", True), Spalte("25w", "E", "D", True),
        Spalte("22m", "F", "G", True), Spalte("23w", "H", "G", True),
        Spalte("20m", "I", "J", True), Spalte("21w", "K", "J", True),
        Spalte("30m", "L", "M", True), Spalte("31w", "N", "M", True),
        Spalte("42m", "O", "P", True), Spalte("43w", "Q", "P", True),
        Spalte("40m", "R", "S", True), Spalte("41w", "T", "S", True),
    ]


def spalten_lichtschiessen() -> list[Spalte]:
    return [
        Spalte("26m", "C", "D", True), Spalte("27w", "E", "D", True),
        Spalte("24m", "F", "G", True), Spalte("25w", "H", "G", True),
        Spalte("22m", "I", "J", True), Spalte("23w", "K", "J", True),
        Spalte("20m", "L", "M", True), Spalte("21w", "N", "M", True),
    ]


# ---------------------------------------------------------------------------
# Konverter
# ---------------------------------------------------------------------------

class Konverter:
    def __init__(self, ws, sportjahr: int) -> None:  # noqa: ANN001
        self.b = Blatt(ws)
        self.sportjahr = sportjahr
        self.disziplinen: dict[str, Disziplin] = {}
        self.befunde: list[Befund] = []
        self.sort = 0

    # ----- Hilfen -------------------------------------------------------------

    def befund(self, art: str, z: Zelle | None, disziplin: str, klasse: str, annahme: str) -> None:
        self.befunde.append(Befund(art, z.koord if z else "", z.wert if z else "", z.fuellung if z else "", disziplin, klasse, annahme))

    def pruefe_header(self, row: int, col: str, erwartet: str, block: str) -> None:
        wert = self.b.text(row, col)
        if str(wert).replace(" ", "").lower() != erwartet.replace(" ", "").lower():
            raise SystemExit(f"Blatt-Layout weicht ab ({block}): {col}{row} = {wert!r}, erwartet {erwartet!r}. Blockdefinition im Skript prüfen.")

    def disziplin(self, kennzahl: str, bezeichnung: str, gruppe: str, quelle: str) -> Disziplin:
        if kennzahl not in self.disziplinen:
            self.sort += 10
            self.disziplinen[kennzahl] = Disziplin(kennzahl=kennzahl, bezeichnung=bezeichnung, gruppe=gruppe, sortierung=self.sort, quelle=quelle)
        return self.disziplinen[kennzahl]

    def ziel(self, gruppe: str, nummer: int, geschlecht: str, d: str, z: Zelle, quelle_klasse: str, kontext: str) -> str | None:
        """Löst eine Verweisnummer in eine Klassenreferenz auf: gleiches Geschlecht bevorzugt,
        sonst die einzige Klasse mit dieser Nummer (geschlechtsübergreifender Verweis)."""
        kl = KLASSEN[gruppe]
        same = f"{nummer}{geschlecht}"
        if same == quelle_klasse:
            self.befund("unklar", z, d, quelle_klasse, f"{kontext}: Verweis auf die eigene Klasse ({z.wert}) – vermutlich verrutschte Zeile im Plan. Als eigene Wertung übernommen, bitte prüfen.")
            return "__selbst__"
        if same in kl:
            return same
        kandidaten = [k for k in kl if k.startswith(str(nummer)) and len(k) == len(str(nummer)) + 1 and not k.endswith("x")]
        if len(kandidaten) == 1:
            self.befund("info", z, d, quelle_klasse, f"{kontext}: Verweis auf {nummer} führt in Klasse {kandidaten[0]} (anderes Geschlecht, im Plan üblich).")
            return kandidaten[0]
        self.befund("unklar", z, d, quelle_klasse, f"{kontext}: Verweis auf {nummer}, aber in Gruppe {gruppe} gibt es keine passende Klasse. Kein Startrecht gesetzt.")
        return None

    # ----- Zellen zu Regel ----------------------------------------------------

    def einzel_setzen(self, disz: Disziplin, sp: Spalte, z: Zelle, gruppe: str, prefix: str) -> None:
        art, wert, hinweis = interpretiere(z)
        r = disz.regel(prefix + sp.klasse)
        geschlecht = sp.klasse[-1]
        if art == "keine":
            if z.wert == "" and z.fuellung == "keine" and gruppe not in ("lichtschiessen",):
                self.befund("info", z, disz.kennzahl, sp.klasse, "Einzel: leere Zelle ohne Rotfüllung → kein Startrecht angenommen.")
            return
        if art == "eigen":
            if wert == "M":
                self.befund("unklar", z, disz.kennzahl, sp.klasse, "Einzel-Spalte enthält 'M' → als eigene Einzelwertung übernommen.")
            r.einzel = "eigen"
            if hinweis:
                r.hinweis = hinweis
            return
        if art == "verweis":
            r.einzel_ziel = self.ziel(gruppe, wert, geschlecht, disz.kennzahl, z, sp.klasse, "Einzel")
            if r.einzel_ziel == "__selbst__":
                r.einzel, r.einzel_ziel = "eigen", None
                return
            r.einzel = "verweis" if r.einzel_ziel else "keine"
            if r.einzel_ziel:
                r.einzel_ziel = prefix + r.einzel_ziel
            return
        if art == "ab18":
            r.mindestalter = 18
            r.hinweis = "ab 18 Jahre"
            r.einzel = "__ab18__"  # wird nach dem Block aus Junioren I übernommen
            return
        if art == "geschlecht_verweis":
            n = wert[geschlecht] if geschlecht in wert else None
            if n is None:
                self.befund("unklar", z, disz.kennzahl, sp.klasse, "Einzel: geschlechtsabhängiger Verweis für Klasse ohne Geschlecht.")
                return
            r.einzel_ziel = self.ziel(gruppe, n, geschlecht, disz.kennzahl, z, sp.klasse, "Einzel")
            r.einzel = "verweis" if r.einzel_ziel else "keine"
            if r.einzel_ziel:
                r.einzel_ziel = prefix + r.einzel_ziel
            self.befund("hinweis", z, disz.kennzahl, sp.klasse, f"Einzel: '{z.wert}' → {geschlecht} startet in {r.einzel_ziel}.")
            return
        if art in ("team", "bei_team"):
            # MixTeam-Zeile: Teamlabel in der Einzelspalte steht für die Mannschaft.
            disz.typ = "mixteam"
            r.mannschaft = "verweis"
            r.mannschaft_ziel = wert
            disz.regel(wert).mannschaft = "eigen"
            disz.hat_mannschaft = True
            return
        self.befund("unklar", z, disz.kennzahl, sp.klasse, f"Einzel: unbekannter Inhalt '{z.wert}' → kein Startrecht gesetzt.")

    def mannschaft_setzen(self, disz: Disziplin, sp: Spalte, z: Zelle, gruppe: str, prefix: str, partner: Spalte | None) -> None:
        art, wert, hinweis = interpretiere(z)
        r = disz.regel(prefix + sp.klasse)
        geschlecht = sp.klasse[-1]
        if art == "keine":
            return
        if art == "eigen":
            if wert == "E":
                self.befund("unklar", z, disz.kennzahl, sp.klasse, "Mannschafts-Spalte enthält 'E' → als eigene Mannschaftswertung übernommen.")
            disz.hat_mannschaft = True
            if sp.geteilt and geschlecht == "w" and partner is not None:
                # Gemeinsame Mannschaftsspalte: w-Klasse verweist auf den Pool der m-Klasse.
                r.mannschaft = "verweis"
                r.mannschaft_ziel = prefix + partner.klasse
            else:
                r.mannschaft = "eigen"
            if hinweis and not r.hinweis:
                r.hinweis = hinweis
            return
        if art == "verweis":
            disz.hat_mannschaft = True
            ziel = self.ziel(gruppe, wert, geschlecht, disz.kennzahl, z, sp.klasse, "Mannschaft")
            if ziel == "__selbst__":
                r.mannschaft = "eigen"
            elif ziel:
                r.mannschaft = "verweis"
                r.mannschaft_ziel = prefix + ziel
            return
        if art == "geschlecht_verweis":
            disz.hat_mannschaft = True
            n = wert.get(geschlecht)
            if n is None:
                self.befund("unklar", z, disz.kennzahl, sp.klasse, "Mannschaft: geschlechtsabhängiger Verweis für Klasse ohne Geschlecht.")
                return
            ziel = self.ziel(gruppe, n, geschlecht, disz.kennzahl, z, sp.klasse, "Mannschaft")
            if ziel:
                r.mannschaft = "verweis"
                r.mannschaft_ziel = prefix + ziel
            self.befund("hinweis", z, disz.kennzahl, sp.klasse, f"Mannschaft: '{z.wert}' → {geschlecht} in Pool {ziel}.")
            return
        if art in ("team", "bei_team"):
            disz.typ = "mixteam"
            disz.hat_mannschaft = True
            r.mannschaft = "verweis"
            r.mannschaft_ziel = wert
            disz.regel(wert).mannschaft = "eigen"
            return
        if art == "ab18":
            self.befund("unklar", z, disz.kennzahl, sp.klasse, "Mannschaft: 'ab 18 Jahre' in Mannschaftsspalte → nicht übernommen.")
            return
        self.befund("unklar", z, disz.kennzahl, sp.klasse, f"Mannschaft: unbekannter Inhalt '{z.wert}' → keine Mannschaft gesetzt.")

    def zeile_verarbeiten(self, row: int, kennzahl: str, name: str, gruppe: str, spalten: list[Spalte], hinweis: str, quelle: str, prefix: str = "") -> None:
        disz = self.disziplin(kennzahl, name, gruppe, quelle)
        if hinweis and hinweis not in disz.hinweis:
            disz.hinweis = (disz.hinweis + " / " + hinweis).strip(" /")
        by_klasse = {s.klasse: s for s in spalten}
        for sp in spalten:
            ze = self.b.zelle(row, sp.einzel)
            self.einzel_setzen(disz, sp, ze, gruppe, prefix)
            if sp.mannschaft is not None:
                zm = self.b.zelle(row, sp.mannschaft)
                partner = None
                if sp.geteilt:
                    n = int(sp.klasse[:-1])
                    partner = by_klasse.get(f"{n - 1}m") if sp.klasse.endswith("w") else None
                self.mannschaft_setzen(disz, sp, zm, gruppe, prefix, partner)
        # "ab 18 Jahre": Regel von Junioren I gleichen Geschlechts übernehmen.
        for sp in spalten:
            r = disz.regeln.get(prefix + sp.klasse)
            if r is None or r.einzel != "__ab18__":
                continue
            geschlecht = sp.klasse[-1]
            vorbild = disz.regeln.get(prefix + ("40m" if geschlecht == "m" else "41w"))
            if vorbild is None or vorbild.einzel == "keine":
                r.einzel = "keine"
                self.befund("unklar", self.b.zelle(row, sp.einzel), disz.kennzahl, sp.klasse, "'ab 18 Jahre', aber Junioren I haben kein Einzel-Startrecht → kein Startrecht gesetzt.")
                continue
            r.einzel = vorbild.einzel
            r.einzel_ziel = vorbild.einzel_ziel
            self.befund("hinweis", self.b.zelle(row, sp.einzel), disz.kennzahl, sp.klasse, f"'ab 18 Jahre' → wie Junioren I ({vorbild.einzel}{' → ' + vorbild.einzel_ziel if vorbild.einzel_ziel else ''}) mit Mindestalter 18.")

    # ----- Blöcke -------------------------------------------------------------

    def hauptblock(self) -> None:
        self.pruefe_header(5, "D", "20", "Hauptblock")
        self.pruefe_header(5, "AJ", "19", "Hauptblock")
        self.pruefe_header(3, "R", "Herren I", "Hauptblock")
        spalten = spalten_hauptblock()
        row = 7
        while row < 68:
            kennzahl = self.b.text(row, "B")
            name = self.b.text(row, "A")
            if kennzahl == "":
                row += 1
                continue
            if kennzahl.startswith(UEBERSPRINGEN_PREFIX):
                self.befund("uebersprungen", self.b.zelle(row, "B"), kennzahl, "", f"{name}: nicht im Umfang (Target Sprint, Sommerbiathlon, Laufende Scheibe, Armbrust).")
                row += 1
                continue
            self.zeile_verarbeiten(row, self.kennzahl_norm(kennzahl), name, "freihand", spalten, self.b.text(row, "AL"), f"Hauptblock Zeile {row}")
            row += 1

    def para_block(self) -> None:
        self.pruefe_header(72, "Z", "90", "Para")
        self.pruefe_header(72, "AH", "96", "Para")
        spalten = spalten_para()
        for row in range(74, 90):
            kennzahl = self.b.text(row, "Y")
            name = self.b.text(row, "W")
            if kennzahl == "":
                continue
            kennzahl = self.kennzahl_norm(kennzahl)
            if kennzahl not in self.disziplinen:
                self.befund("hinweis", self.b.zelle(row, "Y"), kennzahl, "", f"Para-Disziplin {name} kommt im Hauptblock nicht vor; Disziplin nur mit Para-Regeln angelegt.")
                self.disziplin(kennzahl, name, "freihand", f"Para Zeile {row}")
            self.zeile_verarbeiten(row, kennzahl, name, "para", spalten, self.b.text(row, "AJ"), f"Para Zeile {row}", prefix="para:")

    def auflage_block(self) -> None:
        self.pruefe_header(94, "D", "50", "Auflage")
        self.pruefe_header(94, "AD", "81", "Auflage")
        spalten = spalten_auflage()
        for row in range(96, 112):
            kennzahl = self.b.text(row, "B")
            name = self.b.text(row, "A")
            if kennzahl == "":
                continue
            if kennzahl.startswith(UEBERSPRINGEN_PREFIX):
                self.befund("uebersprungen", self.b.zelle(row, "B"), kennzahl, "", f"{name}: nicht im Umfang (Armbrust).")
                continue
            self.zeile_verarbeiten(row, self.kennzahl_norm(kennzahl), name, "auflage", spalten, self.b.text(row, "AF"), f"Auflage Zeile {row}")

    def fitasc_block(self) -> None:
        self.pruefe_header(118, "D", "68", "FITASC")
        self.pruefe_header(118, "T", "66", "FITASC")
        spalten = spalten_fitasc()
        for row in range(120, 124):
            kennzahl = self.b.text(row, "B")
            name = self.b.text(row, "A")
            if kennzahl == "":
                continue
            self.zeile_verarbeiten(row, self.kennzahl_norm(kennzahl), name, "fitasc", spalten, self.b.text(row, "X"), f"FITASC Zeile {row}")

    def blasrohr_block(self) -> None:
        self.pruefe_header(128, "C", "10", "Blasrohr Erwachsene")
        self.pruefe_header(134, "C", "24", "Blasrohr Jugend (Zeile 'Alter' enthält die Klassennummern)")
        for row in (130,):
            kennzahl = self.b.text(row, "B")
            if kennzahl == "1.12":
                self.befund("hinweis", self.b.zelle(row, "B"), "12.10", "", "Blasrohr steht im Plan unter 1.12; korrigiert auf 12.10.")
            self.zeile_verarbeiten(row, "12.10", "Blasrohrsport", "blasrohr", spalten_blasrohr_erwachsene(), "", f"Blasrohr Zeile {row}")
        for row in (136,):
            self.befund("hinweis", self.b.zelle(134, "A"), "12.10", "", "Blasrohr-Jugendblock: Zeilen 'Alter' und 'Klasse' sind vertauscht; Klassennummern aus Zeile 134 verwendet.")
            self.zeile_verarbeiten(row, "12.10", "Blasrohrsport", "blasrohr", spalten_blasrohr_jugend(), "", f"Blasrohr Zeile {row}")

    def lichtschiessen_block(self) -> None:
        # Klassennummern stehen geteilt in zwei Zellen (C143=2, D143=6 → 26).
        def nummer(row: int, c1: str, c2: str | None) -> str:
            a = self.b.text(row, c1)
            b = self.b.text(row, c2) if c2 else ""
            return (a + b).replace(".0", "")
        if nummer(143, "C", "D") != "26" or nummer(143, "E", None) != "27" or nummer(143, "L", "M") != "20":
            raise SystemExit("Lichtschießen-Block: Klassennummern in Zeile 143 nicht wie erwartet (26/27 … 20/21).")
        self.befund("hinweis", self.b.zelle(143, "C"), "11.x", "", "Lichtschießen: Klassennummern 26/24/22/20 stehen auf je zwei Zellen verteilt; zusammengesetzt.")
        spalten = spalten_lichtschiessen()
        for row in range(144, 152):
            kennzahl = self.b.text(row, "B")
            name = self.b.text(row, "A")
            if kennzahl == "":
                continue
            self.zeile_verarbeiten(row, self.kennzahl_norm(kennzahl), name, "lichtschiessen", spalten, self.b.text(row, "O"), f"Lichtschießen Zeile {row}")

    @staticmethod
    def kennzahl_norm(k: str) -> str:
        k = k.strip().rstrip(".")
        k = re.sub(r"\s+", " ", k)
        return k

    # ----- Nachbearbeitung ----------------------------------------------------

    def nachbearbeiten(self) -> None:
        for d in self.disziplinen.values():
            # MixTeams: keine Einzelwertung, Teamklassen nur wenn referenziert.
            if d.typ == "mixteam":
                for r in d.regeln.values():
                    if r.einzel != "keine" and not r.klasse.endswith("x"):
                        self.befund("hinweis", None, d.kennzahl, r.klasse, "MixTeam: Einzelwertung entfernt.")
                    if not r.klasse.endswith("x"):
                        r.einzel = "keine"
                        r.einzel_ziel = None
            # Verweise auf Klassen ohne eigene Wertung prüfen (Warnung, kein Abbruch).
            for r in d.regeln.values():
                for feld, modus in (("einzel_ziel", "einzel"), ("mannschaft_ziel", "mannschaft")):
                    ziel = getattr(r, feld)
                    if not ziel:
                        continue
                    zr = d.regeln.get(ziel)
                    if zr is None or getattr(zr, modus) != "eigen":
                        self.befund("hinweis", None, d.kennzahl, r.klasse, f"{modus.capitalize()}: Verweis auf {ziel}, aber {ziel} hat dort keine eigene Wertung (Kette oder Lücke). Die Engine löst Ketten auf; bitte gegen das PDF prüfen.")
            # Blasrohr/Lichtschießen ohne Mannschaft: Größe 0.
            if not d.hat_mannschaft:
                self.befund("info", None, d.kennzahl, "", "Keine Mannschaftswertung in dieser Disziplin (Mannschaftsgröße 0).")

    # ----- Bogen aus CSV ----------------------------------------------------------

    def bogen_csv(self, pfad: str) -> None:
        with open(pfad, newline="", encoding="utf-8-sig") as fh:
            reader = csv.DictReader(fh, delimiter=";")
            klassen = [c for c in (reader.fieldnames or []) if re.match(r"^\d{2}[mw]$", c)]
            for zeile in reader:
                kennzahl = (zeile.get("kennzahl") or "").strip()
                name = (zeile.get("bezeichnung") or "").strip()
                if kennzahl == "":
                    continue
                d = self.disziplin(kennzahl, name, "bogen", f"Bogen-CSV")
                d.typ = "bogen"
                d.hinweis = (zeile.get("hinweis") or "").strip()
                for k in klassen:
                    z = Zelle(koord=f"{kennzahl}/{k}", wert=(zeile.get(k) or "").strip(), fuellung="keine")
                    sp = Spalte(k, "", None)
                    if k not in KLASSEN["bogen"]:
                        self.befund("unklar", z, kennzahl, k, "Bogen-CSV: unbekannte Klasse.")
                        continue
                    self.einzel_setzen(d, sp, z, "bogen", "")
                    if "*" in z.wert:
                        d.regel(k).hinweis = "keine DM / DSB-Ausschreibung beachten"

    # ----- Ausgabe -----------------------------------------------------------------

    def dokument(self, stand: str) -> dict[str, Any]:
        disziplinen = sorted(self.disziplinen.values(), key=lambda d: (d.gruppe != "bogen", self.sortkey(d.kennzahl)))
        for i, d in enumerate(disziplinen):
            d.sortierung = (i + 1) * 10
        return {
            "format": "kmm-regeltabelle",
            "version": 1,
            "sportjahr": self.sportjahr,
            "stand": stand,
            "disziplinen": [d.to_dict() for d in disziplinen],
        }

    @staticmethod
    def sortkey(kennzahl: str) -> tuple:
        teile = re.split(r"[.\s]+", kennzahl)
        out: list = []
        for t in teile:
            m = re.match(r"^(\d+)(.*)$", t)
            out.append((int(m.group(1)), m.group(2)) if m else (9999, t))
        return tuple(out)

    def validiere(self) -> None:
        """Alle Klassenreferenzen müssen im Klassensatz existieren."""
        for d in self.disziplinen.values():
            for r in d.regeln.values():
                for ref in (r.klasse, r.einzel_ziel, r.mannschaft_ziel):
                    if not ref:
                        continue
                    gruppe, _, kl = ref.rpartition(":")
                    gruppe = gruppe or d.gruppe
                    if kl not in KLASSEN.get(gruppe, {}):
                        self.befund("unklar", None, d.kennzahl, r.klasse, f"Referenz {ref} existiert nicht im Klassensatz {gruppe}.")


# ---------------------------------------------------------------------------
# Berichte
# ---------------------------------------------------------------------------

def schreibe_bericht(pfad: str, k: Konverter, quelle: str) -> None:
    unklar = [b for b in k.befunde if b.art == "unklar"]
    hinweise = [b for b in k.befunde if b.art == "hinweis"]
    infos = [b for b in k.befunde if b.art == "info"]
    uebersprungen = [b for b in k.befunde if b.art == "uebersprungen"]
    gruppen: dict[str, int] = {}
    regeln = 0
    for d in k.disziplinen.values():
        gruppen[d.gruppe] = gruppen.get(d.gruppe, 0) + 1
        regeln += sum(1 for r in d.regeln.values() if not r.leer())

    def tabelle(liste: list[Befund]) -> str:
        if not liste:
            return "_keine_\n"
        out = ["| Zelle | Disziplin | Klasse | Rohwert | Füllung | Annahme |", "|---|---|---|---|---|---|"]
        for b in liste:
            out.append(f"| {b.koord} | {b.disziplin} | {b.klasse} | `{b.wert}` | {b.fuellung} | {b.annahme} |")
        return "\n".join(out) + "\n"

    lines = [
        f"# Prüfbericht Disziplinplan {k.sportjahr}",
        "",
        f"Quelle: `{quelle}`, konvertiert am {date.today().isoformat()}.",
        "",
        "## Übersicht",
        "",
        f"- Disziplinen: {len(k.disziplinen)} ({', '.join(f'{g}: {n}' for g, n in sorted(gruppen.items()))})",
        f"- Regeln (Disziplin × Klasse mit Startrecht): {regeln}",
        f"- **Unklare Zellen (bitte von Hand prüfen): {len(unklar)}**",
        f"- Hinweise (Annahmen des Skripts): {len(hinweise)}",
        f"- Informationen: {len(infos)}",
        f"- Übersprungene Zeilen: {len(uebersprungen)}",
        "",
        "Legende Füllung: rot = kein Startrecht, gruen = Startmöglichkeit nach SpO 0.7.1.1, gelb = olympisch, blau = Änderung gegenüber Vorjahr, orange = Meldung DSB anders.",
        "",
        "## Unklare Zellen",
        "",
        tabelle(unklar),
        "## Hinweise (Annahmen)",
        "",
        tabelle(hinweise),
        "## Übersprungene Zeilen",
        "",
        tabelle(uebersprungen),
        "## Informationen",
        "",
        tabelle(infos),
    ]
    with open(pfad, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))


def schreibe_uebersicht(pfad: str, k: Konverter) -> None:
    """Kompakte Matrix je Disziplin zum Abgleich mit dem PDF."""
    lines = [f"# Regeltabelle {k.sportjahr} – Übersicht zum Abgleich", ""]
    for d in sorted(k.disziplinen.values(), key=lambda d: d.sortierung):
        lines.append(f"## {d.kennzahl} {d.bezeichnung}  ({d.gruppe}, {d.typ}{', ohne Mannschaft' if not d.hat_mannschaft else ''})")
        if d.hinweis:
            lines.append(f"_{d.hinweis}_")
        lines.append("")
        lines.append("| Klasse | Einzel | Mannschaft | Zusatz |")
        lines.append("|---|---|---|---|")
        for r in d.regeln.values():
            if r.leer():
                continue
            gruppe, _, kl = r.klasse.rpartition(":")
            name = KLASSEN.get(gruppe or d.gruppe, {}).get(kl, "?")
            e = "–" if r.einzel == "keine" else ("E" if r.einzel == "eigen" else f"→ {r.einzel_ziel}")
            m = "–" if r.mannschaft == "keine" else ("M" if r.mannschaft == "eigen" else f"→ {r.mannschaft_ziel}")
            zusatz = " ".join(x for x in [f"ab {r.mindestalter}" if r.mindestalter and str(r.mindestalter) not in r.hinweis else "", r.hinweis] if x)
            lines.append(f"| {r.klasse} {name} | {e} | {m} | {zusatz} |")
        lines.append("")
    with open(pfad, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))


def schreibe_bogen_vorlage(pfad: str) -> None:
    klassen = ["10m", "11w", "12m", "13w", "14m", "15w", "20m", "21w", "22m", "23w", "24m", "25w", "30m", "31w", "40m", "41w"]
    disziplinen = [
        ("6.10", "Recurve WA im Freien"), ("6.15", "Compound WA im Freien"), ("6.16", "Blankbogen WA im Freien"),
        ("6.20", "Recurve Halle"), ("6.25", "Compound Halle"), ("6.26", "Blankbogen Halle"),
        ("6.30", "Recurve Feldbogen"), ("6.40", "Blankbogen Feldbogen"), ("6.50", "Compound Feldbogen"),
        ("6.60", "Recurve 3D"), ("6.65", "Compound 3D"), ("6.66", "Blankbogen 3D"), ("6.67", "Langbogen 3D"), ("6.68", "Traditionell 3D"),
    ]
    with open(pfad, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["kennzahl", "bezeichnung", *klassen, "hinweis"])
        for kz, name in disziplinen:
            w.writerow([kz, name, *([""] * len(klassen)), ""])


# ---------------------------------------------------------------------------

def main() -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("xlsx")
    p.add_argument("--sportjahr", type=int, default=2026)
    p.add_argument("--json", default="regeltabelle.json")
    p.add_argument("--bericht", default="pruefbericht.md")
    p.add_argument("--uebersicht", default=None)
    p.add_argument("--bogen-csv", default=None, help="Ausgefüllte Bogen-CSV (aus der Vorlage) einmischen")
    p.add_argument("--bogen-vorlage", default=None, help="Nur die leere Bogen-CSV-Vorlage schreiben und beenden")
    args = p.parse_args()

    if args.bogen_vorlage:
        schreibe_bogen_vorlage(args.bogen_vorlage)
        print(f"Bogen-Vorlage geschrieben: {args.bogen_vorlage}")
        return 0

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    ws = wb.worksheets[0]
    stand = ws["AH1"].value or ""
    k = Konverter(ws, args.sportjahr)
    k.hauptblock()
    k.para_block()
    k.auflage_block()
    k.fitasc_block()
    k.blasrohr_block()
    k.lichtschiessen_block()
    if args.bogen_csv:
        k.bogen_csv(args.bogen_csv)
    k.nachbearbeiten()
    k.validiere()

    doc = k.dokument(f"NSSV 01A1 {stand}, konvertiert am {date.today().isoformat()}")
    with open(args.json, "w", encoding="utf-8") as fh:
        json.dump(doc, fh, ensure_ascii=False, indent=2)
    schreibe_bericht(args.bericht, k, args.xlsx)
    if args.uebersicht:
        schreibe_uebersicht(args.uebersicht, k)

    unklar = sum(1 for b in k.befunde if b.art == "unklar")
    print(f"{len(k.disziplinen)} Disziplinen, {sum(len([r for r in d.regeln.values() if not r.leer()]) for d in k.disziplinen.values())} Regeln → {args.json}")
    print(f"Prüfbericht: {args.bericht} ({unklar} unklare Zellen)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
